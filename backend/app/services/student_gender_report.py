"""Talabgorlar — viloyat × (sana, smena, jins) kesimidagi hisobot ("Gender info").

Qatorlar — viloyatlar, ustunlar — tanlangan test va sana oralig'idagi har bir
(sana, smena) juftligi uchun ikkitadan: Erkak va Ayol. Masalan:
"12-oktabr 1-smena Erkak", "12-oktabr 1-smena Ayol", "12-oktabr 2-smena Erkak", ...

Jins `student_ps_data.gender_id → gender.key` dan olinadi (1 = erkak, 2 = ayol).
Ma'lumot bo'lmasa yoki `key=0` (noma'lum) bo'lsa — loaderdagi kabi JShShIR'ning
birinchi raqamidan aniqlanadi: 3/5 → erkak, 4/6 → ayol. Shunda ham aniqlanmaganlar
ustunlarga kirmaydi, alohida `unknown` sonida qaytadi — jami soni bilan ro'yxat
soni farqi shu bilan tushuntiriladi.

Filtrlar `GET /students` bilan bir xil (`_student_filter_conditions`), shuning
uchun hisobot jadvalda ko'rinayotgan filtr natijasiga mos keladi.
"""

from __future__ import annotations

from dataclasses import dataclass, field
from datetime import date, datetime
from io import BytesIO

from sqlalchemy import case, func, select
from sqlalchemy.orm import Session

from app.crud.student import _student_filter_conditions
from app.models.gender import Gender
from app.models.region import Region
from app.models.smena import Smena
from app.models.student import Student
from app.models.student_ps_data import StudentPsData
from app.models.test import Test
from app.models.test_session import TestSession
from app.models.test_session_smena import TestSessionSmena
from app.models.zone import Zone

MALE = 1
FEMALE = 2
UNKNOWN = 0

_MONTHS_UZ = (
    "yanvar", "fevral", "mart", "aprel", "may", "iyun",
    "iyul", "avgust", "sentabr", "oktabr", "noyabr", "dekabr",
)
GENDER_LABELS = {MALE: "Erkak", FEMALE: "Ayol"}


def format_day_uz(d: date, with_year: bool = False) -> str:
    """`12-oktabr` (yoki oraliq bir necha yilga tushsa `12-oktabr 2026`)."""
    text = f"{d.day}-{_MONTHS_UZ[d.month - 1]}"
    return f"{text} {d.year}" if with_year else text


def smena_label(number: int | None, name: str | None) -> str:
    if number and number > 0:
        return f"{number}-smena"
    return (name or "Smena").strip()


@dataclass
class GenderReportColumn:
    day: date
    smena_number: int
    smena_name: str
    label: str  # "12-oktabr 1-smena"


@dataclass
class GenderReportRow:
    region_id: int
    region_name: str
    # columns tartibida: [[erkak, ayol], ...]
    cells: list[list[int]]
    male: int = 0
    female: int = 0
    unknown: int = 0


@dataclass
class GenderReport:
    test_id: int
    test_name: str | None
    date_from: date
    date_to: date
    columns: list[GenderReportColumn] = field(default_factory=list)
    rows: list[GenderReportRow] = field(default_factory=list)
    # Jami qator — columns tartibida [[erkak, ayol], ...]
    totals: list[list[int]] = field(default_factory=list)
    male: int = 0
    female: int = 0
    unknown: int = 0


def _gender_key_expr():
    """Jins kaliti: avval `gender.key` (1/2), bo'lmasa JShShIR 1-raqami."""
    first_digit = func.substr(Student.imei, 1, 1)
    return case(
        (Gender.key.in_((MALE, FEMALE)), Gender.key),
        (first_digit.in_(("3", "5")), MALE),
        (first_digit.in_(("4", "6")), FEMALE),
        else_=UNKNOWN,
    )


def collect_gender_report(
    db: Session,
    *,
    test_id: int,
    date_from: date,
    date_to: date,
    **filters,
) -> GenderReport:
    """Hisobotni yig'adi. `filters` — `_student_filter_conditions` kalitlari
    (`region_id` qamrovi chaqiruvchida allaqachon qo'llangan bo'lishi kerak)."""
    conditions = _student_filter_conditions(
        test_id=test_id,
        e_date_from=date_from.isoformat(),
        e_date_to=date_to.isoformat(),
        **filters,
    )

    # Ichki so'rov: har talabgor bitta qator + hisoblangan jins kaliti.
    # GROUP BY tashqi so'rovda subquery ustunlari bo'yicha — CASE ifodasini
    # GROUP BY'da takrorlash Postgres'da bind-parametrlar tufayli xato beradi.
    inner = (
        select(
            Student.id.label("sid"),
            Region.id.label("region_id"),
            Region.name.label("region_name"),
            Region.number.label("region_number"),
            Student.e_date.label("e_date"),
            Smena.number.label("smena_number"),
            Smena.name.label("smena_name"),
            _gender_key_expr().label("g"),
        )
        .select_from(Student)
        .join(Zone, Student.zone_id == Zone.id)
        .join(Region, Zone.region_id == Region.id)
        .join(TestSessionSmena, Student.session_smena_id == TestSessionSmena.id)
        .join(Smena, TestSessionSmena.test_smena_id == Smena.id)
        .join(TestSession, TestSessionSmena.test_session_id == TestSession.id)
        .join(Test, TestSession.test_id == Test.id)
        .outerjoin(StudentPsData, StudentPsData.student_id == Student.id)
        .outerjoin(Gender, Gender.id == StudentPsData.gender_id)
    )
    for cond in conditions:
        inner = inner.where(cond)
    sub = inner.subquery()

    stmt = select(
        sub.c.region_id,
        sub.c.region_name,
        sub.c.region_number,
        sub.c.e_date,
        sub.c.smena_number,
        sub.c.smena_name,
        sub.c.g,
        func.count(func.distinct(sub.c.sid)),
    ).group_by(
        sub.c.region_id,
        sub.c.region_name,
        sub.c.region_number,
        sub.c.e_date,
        sub.c.smena_number,
        sub.c.smena_name,
        sub.c.g,
    )
    rows = db.execute(stmt).all()

    report = GenderReport(
        test_id=test_id,
        test_name=db.execute(select(Test.name).where(Test.id == test_id)).scalar(),
        date_from=date_from,
        date_to=date_to,
    )
    if not rows:
        return report

    # Ustunlar: (sana, smena raqami) bo'yicha tartiblangan.
    col_keys: dict[tuple[date, int], str] = {}
    regions: dict[int, tuple[int, str]] = {}
    counts: dict[tuple[int, date, int, int], int] = {}
    for region_id, region_name, region_number, e_date, sm_num, sm_name, g, cnt in rows:
        day = e_date if isinstance(e_date, date) else date.fromisoformat(str(e_date))
        sm_num = int(sm_num or 0)
        g = int(g or 0)
        col_keys.setdefault((day, sm_num), sm_name or "")
        regions[int(region_id)] = (int(region_number or 0), region_name or "")
        key = (int(region_id), day, sm_num, g)
        counts[key] = counts.get(key, 0) + int(cnt)

    with_year = date_from.year != date_to.year
    ordered_cols = sorted(col_keys)
    report.columns = [
        GenderReportColumn(
            day=day,
            smena_number=sm_num,
            smena_name=col_keys[(day, sm_num)],
            label=f"{format_day_uz(day, with_year)} {smena_label(sm_num, col_keys[(day, sm_num)])}",
        )
        for day, sm_num in ordered_cols
    ]
    report.totals = [[0, 0] for _ in ordered_cols]

    for region_id, (_, region_name) in sorted(regions.items(), key=lambda kv: (kv[1][0], kv[1][1])):
        row = GenderReportRow(region_id=region_id, region_name=region_name, cells=[])
        for idx, (day, sm_num) in enumerate(ordered_cols):
            m = counts.get((region_id, day, sm_num, MALE), 0)
            f = counts.get((region_id, day, sm_num, FEMALE), 0)
            row.unknown += counts.get((region_id, day, sm_num, UNKNOWN), 0)
            row.cells.append([m, f])
            row.male += m
            row.female += f
            report.totals[idx][0] += m
            report.totals[idx][1] += f
        report.rows.append(row)
        report.male += row.male
        report.female += row.female
        report.unknown += row.unknown

    return report


# ============================ Excel (.xlsx) ============================

_BRAND = "1E88E5"
_BRAND_DARK = "1565C0"
_HEADER_TEXT = "FFFFFF"
_ZEBRA = "EEF4FB"
_BORDER = "D6E2F0"
_TOTAL_FILL = "DCE8F7"


def build_gender_report_xlsx(report: GenderReport, generated_at: datetime) -> bytes:
    """Ustun nomlari tekis: "12-oktabr 1-smena Erkak", "12-oktabr 1-smena Ayol", ...
    oxirida "Jami Erkak", "Jami Ayol", "Jami"; pastda "Jami" qatori."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Jins bo'yicha"

    headers = ["№", "Viloyat"]
    for col in report.columns:
        headers.append(f"{col.label} {GENDER_LABELS[MALE]}")
        headers.append(f"{col.label} {GENDER_LABELS[FEMALE]}")
    headers += ["Jami Erkak", "Jami Ayol", "Jami"]
    ncols = len(headers)

    thin = Side(style="thin", color=_BORDER)
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left = Alignment(horizontal="left", vertical="center", wrap_text=True)

    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncols)
    title = ws.cell(row=1, column=1, value="TALABGORLAR — JINS BO'YICHA HISOBOT")
    title.font = Font(bold=True, size=14, color=_HEADER_TEXT)
    title.alignment = Alignment(horizontal="center", vertical="center")
    title.fill = PatternFill("solid", fgColor=_BRAND_DARK)
    ws.row_dimensions[1].height = 26

    period = (
        f"{report.date_from.strftime('%d.%m.%Y')} — {report.date_to.strftime('%d.%m.%Y')}"
    )
    meta_text = (
        f"Test: {report.test_name or report.test_id}    |    Sana: {period}    |    "
        f"Erkak: {report.male}, Ayol: {report.female}, Jami: {report.male + report.female}"
    )
    if report.unknown:
        meta_text += f"    |    Jinsi aniqlanmagan: {report.unknown}"
    meta_text += f"    |    Yuklab olindi: {generated_at.strftime('%d.%m.%Y %H:%M')}"
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=ncols)
    meta = ws.cell(row=2, column=1, value=meta_text)
    meta.font = Font(italic=True, size=10, color="5B6B7B")
    meta.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[2].height = 18

    header_row = 3
    header_fill = PatternFill("solid", fgColor=_BRAND)
    header_font = Font(bold=True, size=10, color=_HEADER_TEXT)
    for col_idx, label in enumerate(headers, start=1):
        c = ws.cell(row=header_row, column=col_idx, value=label)
        c.font = header_font
        c.fill = header_fill
        c.alignment = center
        c.border = border
    ws.row_dimensions[header_row].height = 46

    def _write_row(excel_row: int, values: list, *, fill: str | None, bold: bool) -> None:
        for col_idx, val in enumerate(values, start=1):
            c = ws.cell(row=excel_row, column=col_idx, value=val)
            c.border = border
            c.alignment = left if col_idx == 2 else center
            if bold:
                c.font = Font(bold=True)
            if fill:
                c.fill = PatternFill("solid", fgColor=fill)

    for i, row in enumerate(report.rows):
        values: list = [i + 1, row.region_name]
        for m, f in row.cells:
            values += [m, f]
        values += [row.male, row.female, row.male + row.female]
        _write_row(header_row + 1 + i, values, fill=_ZEBRA if i % 2 else None, bold=False)

    total_row = header_row + 1 + len(report.rows)
    total_values: list = ["", "Jami"]
    for m, f in report.totals:
        total_values += [m, f]
    total_values += [report.male, report.female, report.male + report.female]
    _write_row(total_row, total_values, fill=_TOTAL_FILL, bold=True)

    ws.column_dimensions["A"].width = 5
    ws.column_dimensions["B"].width = 26
    for col_idx in range(3, ncols + 1):
        ws.column_dimensions[get_column_letter(col_idx)].width = 13
    ws.freeze_panes = ws.cell(row=header_row + 1, column=3)
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
